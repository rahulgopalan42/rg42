# Audit step 1: extract verified templates from the Word docs and inventory master-sheet markings.
# Usage: python3.13 -I audit_extract.py <docs_dir> <master.xlsx>
# Writes audit_docs.json (one row per email section) and audit_marks.json (per-tab raw view).
import collections, glob, json, os, re, sys
import docx, openpyxl

DOCS, MASTER = sys.argv[1], sys.argv[2]
OUT = os.path.dirname(os.path.abspath(__file__))
HDR = re.compile(r'^(\d+)\.\s+(.+?)\s+\|\s+group:\s*(.+)$')
HOOK = re.compile(r'looking for an? (.+?), and (.+?) was identified as a potential fit')

rows = []
for path in sorted(glob.glob(f'{DOCS}/*.docx')):
    doc = os.path.basename(path)[:-5]
    if doc == 'W10b':
        continue                     # byte-identical duplicate of W10
    title = None
    cur = None
    for p in docx.Document(path).paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if title is None:
            title = t
        m = HDR.match(t)
        if m:
            cur = {'doc': doc, 'n': int(m.group(1)), 'header': m.group(2).strip(), 'group': m.group(3).strip(),
                   'opener': False, 'hook_brand': None, 'designation': None, 'greeting': False}
            rows.append(cur)
            continue
        if cur is None:
            continue
        if t.startswith('Reaching out from SPORTFIVE'):
            cur['opener'] = True
        if t.startswith('Hi '):
            cur['greeting'] = True
        h = HOOK.search(t)
        if h and not cur['hook_brand']:
            cur['designation'], cur['hook_brand'] = h.group(1), h.group(2)
    print(f'{doc}: {sum(1 for r in rows if r["doc"] == doc)} sections | {title[:60]}')

bad = [r for r in rows if not (r['opener'] and r['hook_brand'] and r['greeting'])]
mism = [r for r in rows if r['hook_brand'] and re.sub(r'[^a-z0-9]', '', r['hook_brand'].lower()) !=
        re.sub(r'[^a-z0-9]', '', r['header'].lower())]
print(f'TOTAL sections {len(rows)} | incomplete emails {len(bad)} | header/hook name differs {len(mism)}')
for r in bad[:10]:
    print('  INCOMPLETE', r['doc'], r['n'], r['header'], {k: r[k] for k in ('opener', 'greeting', 'hook_brand')})
for r in mism[:15]:
    print('  DIFF', r['doc'], r['n'], repr(r['header']), '->', repr(r['hook_brand']))
json.dump(rows, open(f'{OUT}/audit_docs.json', 'w'), indent=0)

# ---- master markings inventory ----
wb = openpyxl.load_workbook(MASTER)          # full mode: needs fills and comments
marks = {}
for ws in wb.worksheets:
    if ws.title == 'OVERVIEW':
        continue
    hdr = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    tokens, fills, fonts = collections.Counter(), collections.Counter(), collections.Counter()
    comments = 0
    sample = []
    for r in ws.iter_rows(min_row=2):
        vals = [c.value for c in r]
        if all(v in (None, '') for v in vals):
            continue
        if len(sample) < 2:
            sample.append([str(v)[:18] if v is not None else '' for v in vals[:6]])
        for c in r[:5]:
            v = c.value
            if isinstance(v, str) and len(v.strip()) <= 4 and not v.strip().isdigit():
                tokens[f'{c.column_letter}:{v.strip()}'] += 1
            f = c.fill
            if f is not None and f.fill_type and f.fgColor is not None and f.fgColor.rgb not in (None, '00000000', 'FFFFFFFF'):
                fills[f'{c.column_letter}:{f.fgColor.rgb if isinstance(f.fgColor.rgb, str) else "theme"}'] += 1
            if c.font is not None and (c.font.strike or (c.font.color is not None and isinstance(c.font.color.rgb, str)
                                                          and c.font.color.rgb not in ('FF000000',))):
                fonts[f'{c.column_letter}:{"strike" if c.font.strike else c.font.color.rgb}'] += 1
            if c.comment:
                comments += 1
    marks[ws.title] = {'header': hdr[:9], 'sample': sample, 'tokens': dict(tokens.most_common(8)),
                       'fills': dict(fills.most_common(5)), 'fonts': dict(fonts.most_common(4)), 'comments': comments}
json.dump(marks, open(f'{OUT}/audit_marks.json', 'w'), indent=0)
for t, m in marks.items():
    print(f'== {t[:32]:32s} hdr={m["header"][:4]} tok={m["tokens"]} fill={m["fills"]} font={m["fonts"]} cmt={m["comments"]}')
