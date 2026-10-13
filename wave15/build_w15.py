"""Wave 15 coverage sweep compiler.

Merges every Clay / Apollo pull in wave15/, removes anything already in the
user's files (known_index.json), auto-screens F1 sponsors, sanctions, gov /
non-profit, under-51 staff and suspect records, maps each brand to the 56
master sectors (plus X tabs for brands outside them) and pre-ranks it.

Usage: python3 -I wave15/build_w15.py <out.xlsx>
"""
import csv, glob, json, math, re, sys, unicodedata
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

W = '/home/user/rg42/wave15'
R = '/home/user/rg42/reports'
csv.field_size_limit(10**8)


def strict(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'\(.*?\)', '', s)
    s = s.replace('&', ' and ').replace('.', ' ').replace('/', ' ').replace(',', ' ')
    s = re.sub(r'\b(inc|incorporated|ltd|limited|plc|corp|corporation|co|company|group|holdings|holding|ag|sa|spa|se|nv|llc|lp|l p|gmbh|as|the|sas|bv|ab|pte|pty|pvt|international|bank|financial|services)\b', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s)


def dom(d):
    d = str(d or '').lower().strip()
    d = re.sub(r'^https?://', '', d); d = re.sub(r'^www\.', '', d)
    d = d.split('/')[0].split(' ')[0]
    return d if '.' in d and len(d) > 4 else ''


# ---------------------------------------------------------------- sectors
SECTORS = ['01 Banking & Insurance', '02 Energy & Industrial', '03 Luxury & Beauty', '04 Food & Beverage',
           '05 Supercomputing', '06 UCaaS & CPaaS', '07 Marketing & Sales', '08 Devtools & Integration',
           '09 Fintech', '10 ERP & Finance Procurement', '11 Storage Hardware', '12 Trading & Crypto',
           '13 AI Agents', '14 Cyber + IT', '15 Observability', '16 Travel & Hospitality',
           '17 AI Cloud & Data', '18 Services & Other', '19 Gulf Sovereign & Tourism', '20 Staffing & Talent',
           '21 Materials & Mech Parts', '22 SaaS', '23 Machine Tools', '24 Logistics',
           '25 IT Services & Consulting', '26 Horizontal SaaS', '27 Chip Makers', '28 AI Data Services',
           '29 Tech Media & Gaming', '30 Commercial Vehicles', '31 GRC & Risk SW', '32 Defense & Industrial AI',
           '33 Automation & Robotics', '34 DC Power & Cooling', '35 Semiconductor Equipment', '36 IoT & Connectivity',
           '37 EDA', '38 Edge & CDN', '39 Construction & Field Svc', '40 Test & Measurement',
           '41 OT & Physical Security', '42 Health & Wellness', '43 Mobility & Automotive', '44 Grid & Transformers',
           '45 Electronic Parts', '46 Tools & Industrial Supply', '47 Vertical SaaS', '48 Telecom & Space',
           '49 Home & Living', '50 Law Firms & Legal AI', '51 Industrial Gases', '52 Photonics & Optical',
           '53 Consumer & Lifestyle', '54 TIC', '55 Packaging & Food Machinery', '56 Robotics AI']
S = {s[:2]: s for s in SECTORS}

IND = {  # Clay industry -> sector code or X tab
    'Banking': '01', 'Insurance': '01', 'Insurance Agencies and Brokerages': '01', 'Insurance Carriers': '01',
    'Financial Services': '01', 'Investment Banking': '12', 'Capital Markets': '12',
    'Investment Management': 'X Asset Mgmt & PE',
    'Oil and Gas': '02', 'Oil, Gas, and Mining': '02', 'Renewable Energy Power Generation': '02',
    'Renewable Energy Equipment Manufacturing': '02', 'Services for Renewable Energy': '02',
    'Renewables & Environment': '02', 'Electric Power Generation': '02', 'Solar Electric Power Generation': '02',
    'Nuclear Electric Power Generation': '02',
    'Utilities': 'X Utilities & Power', 'Natural Gas Distribution': 'X Utilities & Power',
    'Electric Power Transmission, Control, and Distribution': '44', 'Electrical Equipment Manufacturing': '44',
    'Retail Luxury Goods and Jewelry': '03', 'Luxury Goods and Jewelry': '03', 'Cosmetics': '03',
    'Personal Care Product Manufacturing': '03', 'Fashion Accessories Manufacturing': '03',
    'Food and Beverage Manufacturing': '04', 'Food & Beverages': '04', 'Beverage Manufacturing': '04',
    'Food Production': '04', 'Dairy Product Manufacturing': '04', 'Dairy': '04', 'Baked Goods Manufacturing': '04',
    'Meat Products Manufacturing': '04', 'Wine and Spirits': '04', 'Distilleries': '04', 'Breweries': '04',
    'Food and Beverage Services': '04', 'Restaurants': '04',
    'Food and Beverage Retail': 'X Retail & E-commerce', 'Retail Groceries': 'X Retail & E-commerce',
    'Telecommunications': '48', 'Wireless Services': '48', 'Satellite Telecommunications': '48',
    'Media & Telecommunications': '48', 'Space Research and Technology': '48',
    'Communications Equipment Manufacturing': '36', 'Computer Networking Products': '36', 'Computer Networking': '36',
    'Advertising Services': '07', 'Marketing Services': '07', 'Public Relations and Communications Services': '07',
    'Market Research': '07',
    'Software Development': '22', 'Technology, Information and Internet': '22', 'Internet Publishing': '22',
    'Information Services': '22', 'Mobile Computing Software Products': '22', 'Embedded Software Products': '22',
    'Data Infrastructure and Analytics': '17', 'Business Intelligence Platforms': '17',
    'Computer and Network Security': '14', 'Data Security Software Products': '14',
    'IT Services and IT Consulting': '25', 'Information Technology and Services': '25',
    'Outsourcing and Offshoring Consulting': '25',
    'Hospitality': '16', 'Hotels and Motels': '16', 'Travel Arrangements': '16', 'Leisure, Travel & Tourism': '16',
    'Airlines and Aviation': '16',
    'Gambling Facilities and Casinos': 'X Gambling & Lotteries', 'Events Services': 'X Live Events & Ticketing',
    'Business Consulting and Services': '18', 'Professional Services': '18', 'Accounting': '18',
    'Facilities Services': '18', 'Engineering Services': '18', 'Security and Investigations': '41',
    'Law Practice': '50', 'Legal Services': '50',
    'Staffing and Recruiting': '20', 'Human Resources Services': '20',
    'Chemical Manufacturing': '21', 'Chemical Raw Materials Manufacturing': '21',
    'Paint, Coating, and Adhesive Manufacturing': '21', 'Plastics Manufacturing': '21',
    'Glass, Ceramics and Concrete Manufacturing': '21', 'Building Materials': '21', 'Primary Metal Manufacturing': '21',
    'Fabricated Metal Products': '21', 'Paper and Forest Product Manufacturing': '21',
    'Mining': 'X Mining & Metals', 'Metal Ore Mining': 'X Mining & Metals',
    'Transportation, Logistics, Supply Chain and Storage': '24', 'Truck Transportation': '24',
    'Freight and Package Transportation': '24', 'Warehousing and Storage': '24', 'Maritime Transportation': '24',
    'Maritime': '24', 'Rail Transportation': '24', 'Ground Passenger Transportation': '24',
    'Semiconductor Manufacturing': '27', 'Semiconductors': '27', 'Renewable Energy Semiconductor Manufacturing': '27',
    'Computer Games': '29', 'Mobile Gaming Apps': '29', 'Entertainment Providers': '29', 'Entertainment': '29',
    'Broadcast Media Production and Distribution': '29', 'Media Production': '29',
    'Online Audio and Video Media': '29', 'Movies, Videos and Sound': '29', 'Radio and Television Broadcasting': '29',
    'Spectator Sports': 'X Sports Clubs & Leagues',
    'Motor Vehicle Manufacturing': '43', 'Automotive': '43', 'Motor Vehicle Parts Manufacturing': '43',
    'Retail Motor Vehicles': '43',
    'Defense and Space Manufacturing': '32', 'Defense & Space': '32',
    'Aviation & Aerospace': 'X Aviation & Aerospace', 'Aviation and Aerospace Component Manufacturing': 'X Aviation & Aerospace',
    'Automation Machinery Manufacturing': '33', 'Industrial Automation': '33',
    'Machinery Manufacturing': '23', 'Industrial Machinery Manufacturing': '23', 'Mechanical or Industrial Engineering': '23',
    'Wholesale': '46',
    'Construction': '39', 'Civil Engineering': '39', 'Building Construction': '39', 'Specialty Trade Contractors': '39',
    'Residential Building Construction': '39', 'Nonresidential Building Construction': '39',
    'Architecture and Planning': '39',
    'Real Estate': 'X Real Estate', 'Leasing Non-residential Real Estate': 'X Real Estate',
    'Commercial Real Estate': 'X Real Estate',
    'Hospitals and Health Care': '42', 'Medical Practices': '42', 'Mental Health Care': '42',
    'Home Health Care Services': '42', 'Hospitals': '42', 'Health, Wellness and Fitness': '42',
    'Medical and Diagnostic Laboratories': '42',
    'Medical Equipment Manufacturing': 'X Medtech', 'Medical Devices': 'X Medtech',
    'Pharmaceutical Manufacturing': 'X Pharma & Biotech', 'Biotechnology Research': 'X Pharma & Biotech',
    'Biotechnology': 'X Pharma & Biotech', 'Retail Pharmacies': 'X Retail & E-commerce',
    'Retail': 'X Retail & E-commerce', 'Online and Mail Order Retail': 'X Retail & E-commerce',
    'Internet Marketplace Platforms': 'X Retail & E-commerce', 'Retail Health and Personal Care Products': 'X Retail & E-commerce',
    'Retail Furniture and Home Furnishings': '49', 'Retail Appliances, Electrical, and Electronic Equipment': 'X Retail & E-commerce',
    'Retail Building Materials and Garden Equipment': 'X Retail & E-commerce', 'Retail Gasoline': '02',
    'Retail Apparel and Fashion': '53', 'Apparel and Fashion': '53', 'Apparel Manufacturing': '53',
    'Footwear Manufacturing': '53', 'Sporting Goods': '53', 'Sporting Goods Manufacturing': '53', 'Consumer Goods': '53',
    'Furniture': '49', 'Furniture and Home Furnishings Manufacturing': '49', 'Household Appliance Manufacturing': '49',
    'Appliances, Electrical, and Electronics Manufacturing': '45', 'Consumer Electronics': 'X Consumer Electronics',
    'Computers and Electronics Manufacturing': '45', 'Computer Hardware Manufacturing': '45', 'Computer Hardware': '45',
    'Measuring and Control Instrument Manufacturing': '40',
    'Packaging and Containers Manufacturing': '55', 'Packaging and Containers': '55',
    'Shipbuilding': 'X Maritime & Shipbuilding', 'Railroad Equipment Manufacturing': 'X Rail & Transit',
    'Higher Education': 'X Education', 'Education': 'X Education', 'E-Learning Providers': 'X Education',
}

NAICS = [  # (prefix, code) - longest prefix wins
    ('52', '09'), ('5242', '01'), ('5241', '01'), ('5221', '01'), ('5239', '12'), ('5231', '12'), ('5259', 'X Asset Mgmt & PE'),
    ('5112', '22'), ('5132', '22'), ('5182', '17'), ('518', '17'), ('5415', '22'), ('541511', '22'), ('541512', '25'),
    ('541519', '25'), ('519', '22'), ('3344', '27'), ('3345', '40'), ('3341', '45'), ('3342', '36'), ('3343', '45'),
    ('336', '43'), ('3364', 'X Aviation & Aerospace'), ('3254', 'X Pharma & Biotech'), ('5417', 'X Pharma & Biotech'),
    ('325', '21'), ('621', '42'), ('622', '42'), ('623', '42'), ('3391', 'X Medtech'), ('454', 'X Retail & E-commerce'),
    ('455', 'X Retail & E-commerce'), ('44', 'X Retail & E-commerce'), ('45', 'X Retail & E-commerce'),
    ('72', '16'), ('721', '16'), ('48', '24'), ('49', '24'), ('481', '16'), ('517', '48'), ('5418', '07'),
    ('5416', '18'), ('5411', '50'), ('5613', '20'), ('23', '39'), ('531', 'X Real Estate'), ('61', 'X Education'),
    ('7132', 'X Gambling & Lotteries'), ('211', '02'), ('213', '02'), ('2211', '02'), ('221', 'X Utilities & Power'),
    ('333', '23'), ('3332', '35'), ('311', '04'), ('312', '04'), ('71', '29'), ('512', '29'), ('516', '29'),
    ('5191', '22'), ('212', 'X Mining & Metals'), ('331', '21'), ('332', '21'), ('335', '44'), ('337', '49'),
    ('3221', '21'), ('3222', '55'), ('326', '21'), ('327', '21'), ('42', '46'), ('5413', '18'), ('5419', '18'),
    ('561', '18'), ('562', 'X Waste & Recycling'), ('111', 'X Agriculture'), ('112', 'X Agriculture'), ('92', 'GOV'),
]

KW = [  # name/domain regex overrides (checked in order, first hit wins)
    (r'supercomput|\bhpc\b', '05'),
    (r'\bucaas\b|\bcpaas\b|ringcentral|vonage|\b8x8\b|twilio|sinch|bandwidth inc|dialpad|zoom video|kaleyra|infobip', '06'),
    (r'devops|gitlab|jfrog|postman|\bdeveloper\b|harness\.io|mulesoft|boomi', '08'),
    (r'\berp\b|procure|spend management|coupa|ivalua|jaggaer|\bbilling\b|accounts payable', '10'),
    (r'storage|netapp|pure storage|\bnas\b|backup', '11'),
    (r'crypto|bitcoin|blockchain|\bcoin|\bbtc\b|\bweb3\b|stablecoin|\bdefi\b|exchange|brokerage|trading|\bmarkets\b', '12'),
    (r'observability|datadog|new relic|dynatrace|grafana|\bapm\b|logging|monitoring', '15'),
    (r'\bgrc\b|compliance|governance|\brisk\b|regtech|\baml\b|\bkyc\b', '31'),
    (r'data ?cent(er|re)|\bcooling\b|liquid cool|\bups\b|vertiv|schneider', '34'),
    (r'lithograph|wafer|deposition|\betch|semiconductor equipment|\bmetrology\b', '35'),
    (r'\biot\b|connectivity|\besim\b|lorawan|\bm2m\b', '36'),
    (r'design automation|\beda\b|synopsys|cadence|siemens eda', '37'),
    (r'\bcdn\b|\bedge\b|akamai|fastly|cloudflare', '38'),
    (r'\btest(ing)? (and|&) measurement|oscillo|keysight|teradyne|\bmeasurement\b|instrument', '40'),
    (r'\bot security|physical security|surveillance|access control|cctv|alarm', '41'),
    (r'transformer|switchgear|\bgrid\b|substation|\bcable[s]?\b', '44'),
    (r'industrial gas|air liquide|\blinde\b|air products|messer|nippon sanso|\bgases\b', '51'),
    (r'photonic|\boptic|\blaser|lidar|\blens', '52'),
    (r'\btic\b|testing,? inspection|certification|bureau veritas|\bsgs\b|intertek|\btuv\b|\btüv\b|dekra|applus|eurofins', '54'),
    (r'robot|humanoid|\bcobot', '56'),
    (r'\btruck|\bbus(es)?\b|trailer|\blorr|tractor|forklift', '30'),
    (r'machine tool|\bcnc\b|\blathe|milling|grinding', '23'),
    (r'\btools?\b|fastener|\bmro\b|industrial supply', '46'),
    (r'packag|bottling|filling machine|food machinery|food processing equipment', '55'),
    (r'\bdefen[cs]e\b|\bmissile|\bmunition|\bdrone|\buav\b|\barmou?r', '32'),
    (r'\bautomation\b|\bplc\b|motion control', '33'),
    (r'\bchip|semiconductor|silicon|\bfab\b|\bfoundry', '27'),
    (r'\blabel(l)?ing\b|annotation|training data|data labeling|scale ai', '28'),
    (r'\bagent(s|ic)?\b|copilot|\bassistant', '13'),
    (r'\bgpu cloud|\bai cloud|\bneocloud|\bai infra|\binference\b', '17'),
    (r'\bpayment|\bpay\b|\bpayroll|\bneobank|\blending\b|\bloan|\bwallet|fintech|\bcard[s]?\b', '09'),
    (r'\blegal\b|\blaw\b|\battorney|\bllp\b', '50'),
    (r'\bpet\b|\bpets\b|petcare|veterinar', 'X Pet Care'),
    (r'watch(es|maker)|jewel', '03'),
]

GULF = {'united arab emirates', 'uae', 'saudi arabia', 'qatar', 'bahrain', 'kuwait', 'oman'}

SANCTION = [  # 1260H / Entity List / UFLPA / OFAC SDN headline names (substring on strict name)
    'huawei', 'hikvision', 'dahua', 'zte', 'smic', 'semiconductormanufacturinginternational', 'dji', 'catl',
    'contemporaryamperex', 'tencent', 'cosco', 'cnooc', 'chinamobile', 'chinatelecom', 'chinaunicom', 'inspur', 'sugon',
    'iflytek', 'sensetime', 'megvii', 'yitu', 'ymtc', 'yangtzememory', 'cxmt', 'changxin', 'avic', 'cssc',
    'norinco', 'hesai', 'quectel', 'bgi', 'autel', 'crrc', 'chinarailwayconstruction', 'chinacommunicationsconstruction',
    'cgn', 'chinageneralnuclear', 'sinochem', 'casc', 'casic', 'cetc', 'chinaelectronicstechnology',
    'gazprom', 'rosneft', 'lukoil', 'sberbank', 'vtb', 'rostec', 'novatek', 'alrosa', 'surgutneftegas', 'transneft',
    'rosatom', 'aeroflot', 'sovcomflot', 'rusal', 'nornickel', 'norilsknickel', 'yandex', 'kaspersky', 'mtsbank',
    'gazprombank', 'alfabank', 'tinkoff', 'promsvyazbank', 'irancell', 'mahanair', 'nioc', 'naftiran',
    'xinjiang', 'zhongtian', 'longi', 'jinkosolar', 'tongwei', 'daqo', 'hoshine', 'esquel',
]

# ---------------------------------------------------------------- inputs
known = json.load(open(f'{W}/known_index.json'))
KN, KD = set(known['names']), set(known['doms'])

sponsors = set()
for row in csv.DictReader(open(f'{R}/f1-sponsor-competitor-map-ring2-2026-08.csv')):
    if row.get('f1_sponsor'): sponsors.add(row['f1_sponsor'])
for row in csv.DictReader(open(f'{R}/f1-sponsor-competitor-map-2026-08.csv')):
    for s in (row.get('f1_sponsors_it_competes_with') or '').split(';'):
        if s.strip(): sponsors.add(s.strip())
    if (row.get('already_in_f1') or '').upper() == 'YES': sponsors.add(row['brand'])
SP = set()
for s in sponsors:
    for part in re.split(r'\(|\)| - ', s):
        n = strict(part)
        if len(n) >= 3: SP.add(n)
ties = json.load(open('/home/user/rg42/wave14/f1_ties_found.json'))
TIES = {strict(k) for k in ties}


def rev_bucket_score(b):
    return {'100B-1T': 100, '10B-100B': 90, '1B-10B': 75, '500M-1B': 62, '200M-500M': 50, '75M-200M': 40}.get(b, 0)


def parse_money(s):
    m = re.match(r'\s*\$?([\d.]+)\s*([KMBT]?)', str(s or ''), re.I)
    if not m: return None
    v = float(m.group(1)); u = m.group(2).upper()
    return v * {'K': 1e3, 'M': 1e6, 'B': 1e9, 'T': 1e12, '': 1}[u]


def money_bucket(v):
    if v is None: return ''
    if v >= 1e11: return '100B-1T'
    if v >= 1e10: return '10B-100B'
    if v >= 1e9: return '1B-10B'
    if v >= 5e8: return '500M-1B'
    if v >= 2e8: return '200M-500M'
    if v >= 7.5e7: return '75M-200M'
    return '<75M'


def emp_int(e):
    try: return int(re.sub(r'[^\d]', '', str(e)) or 0)
    except ValueError: return 0


def size_floor(sz):
    m = re.match(r'([\d,]+)', str(sz or ''))
    return int(m.group(1).replace(',', '')) if m else 0


recs = {}
order = []


def add(r):
    k = r['domain'] or ('n:' + r['key'])
    if not r['key'] and not r['domain']: return
    if k in recs:
        o = recs[k]
        o['sources'] |= r['sources']
        if r['score_rev'] > o['score_rev']:
            for f in ('rev_band', 'score_rev'): o[f] = r[f]
        for f in ('industry', 'naics', 'country', 'type', 'linkedin', 'owner', 'exchange'):
            if not o.get(f) and r.get(f): o[f] = r[f]
        o['emp'] = max(o['emp'], r['emp'])
        o['funded'] = o['funded'] or r['funded']
        return
    recs[k] = r; order.append(k)


for f in sorted(glob.glob(f'{W}/clay_*.tsv')):
    for x in csv.DictReader(open(f, encoding='utf-8', errors='replace'), delimiter='\t'):
        name = (x.get('name') or '').strip()
        if not name: continue
        emp = emp_int(x.get('employee_count')) or size_floor(x.get('size'))
        add(dict(name=name, key=strict(name), domain=dom(x.get('domain')), country=(x.get('country') or '').strip(),
                 type=(x.get('type') or '').strip(), emp=emp, size=x.get('size') or '', rev_band=x.get('annual_revenue') or '',
                 score_rev=rev_bucket_score(x.get('annual_revenue')), industry=x.get('industry') or '', naics='',
                 linkedin=x.get('linkedin_url') or '', owner='', exchange='', funded=False, sources={'Clay'}))

for f in sorted(glob.glob(f'{W}/apollo_*.tsv')):
    for x in csv.DictReader(open(f, encoding='utf-8', errors='replace'), delimiter='\t'):
        name = (x.get('name') or '').strip()
        if not name: continue
        v = parse_money(x.get('organization_revenue_printed'))
        b = money_bucket(v)
        owner = ''
        try:
            o = json.loads(x.get('owned_by_organization') or 'null')
            if isinstance(o, dict): owner = o.get('name') or ''
        except Exception:
            pass
        try: naics = json.loads(x.get('naics_codes') or '[]') or []
        except Exception: naics = []
        add(dict(name=name, key=strict(name), domain=dom(x.get('primary_domain')), country='',
                 type='Public Company' if x.get('publicly_traded_exchange') else '', emp=51, size='51+',
                 rev_band=b, score_rev=rev_bucket_score(b), industry='', naics=','.join(map(str, naics)),
                 linkedin=x.get('linkedin_url') or '', owner=owner, exchange=x.get('publicly_traded_exchange') or '',
                 funded=True, sources={'Apollo'}))


def sector_of(r):
    text = f"{r['name']} {r['domain']}".lower()
    for rx, code in KW:
        if re.search(rx, text):
            # keyword hits only override generic buckets, not clear industry signals
            base = IND.get(r['industry'])
            if base in (None, '22', '25', '18', '45', '21', '01', '23', '46') or code in ('51', '54', '56', '37', '05'):
                return code
            break
    if r['industry'] in IND:
        code = IND[r['industry']]
    else:
        code = None
        best = 0
        for n in r['naics'].split(','):
            for p, c in NAICS:
                if n.startswith(p) and len(p) > best: best, code = len(p), c
        code = code or 'X Unsorted'
    if code in ('16', 'X Real Estate', 'GOV') and r['country'].lower() in GULF:
        code = '19'
    return code


def tabname(code):
    if code in S: return S[code]
    return code[:31]


GEN = set('global general national international american united first china india japan new group energy power '
          'capital digital data systems solutions technologies technology tech services industries industrial health '
          'care medical bank financial auto motor motors security royal premier pacific atlantic north south east '
          'west central euro asia africa star sun green blue red city metro grand great best smart total prime alpha '
          'delta omega hotel airlines fashion universal associated'.split())
GEO = set('usa us america americas canada mexico brasil brazil argentina chile colombia peru uk britain ireland france '
          'germany deutschland italia italy espana spain portugal nederland netherlands belgium benelux nordic nordics '
          'sweden norway denmark finland poland polska czech austria schweiz switzerland suisse turkey turkiye greece '
          'romania hungary russia ukraine israel uae arabia saudi qatar egypt africa nigeria kenya india pakistan '
          'bangladesh china japan korea taiwan hongkong hong kong singapore malaysia indonesia thailand vietnam '
          'philippines australia zealand nz europe emea apac asia latam latin middle east global international inc '
          'ltd llc gmbh sa srl spa pvt private limited co corp company group do de da du di del of the and'.split())
BADCOUNTRY = {'russia', 'russian federation', 'belarus', 'iran', 'north korea', 'syria', 'cuba'}


def toks(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'\(.*?\)', '', s)
    return [t for t in re.split(r'[^a-z0-9]+', s) if t]


def arm_of_known(r):
    t = toks(r['name'])
    if len(t) < 2: return None
    two = ''.join(t[:2])
    if len(t) > 2 and two in KN: return two
    one = t[0]
    if len(one) < 5 or one in GEN or one not in KN: return None
    label = r['domain'].split('.')[0] if r['domain'] else ''
    if one in label or all(x in GEO for x in t[1:]): return one
    return None


def sld(d):
    p = d.split('.')
    if len(p) >= 3 and p[-2] in ('co', 'com', 'net', 'org', 'gov', 'ac') and len(p[-1]) == 2: return p[-3]
    return p[-2] if len(p) >= 2 else d


# F1 teams, power-unit makers, team owners and rights/broadcast holders (and their groups) - always a tie
F1_CORE = {strict(x) for x in (
    'General Motors', 'GM', 'GM Financial', 'Cadillac', 'Ford', 'Ford Motor', 'Ford Motor Company', 'Toyota', 'Honda',
    'Audi', 'Volkswagen', 'Mercedes-Benz', 'Mercedes', 'Daimler', 'Ferrari', 'Renault', 'Alpine', 'Aston Martin',
    'McLaren', 'Red Bull', 'Williams', 'Haas', 'Sauber', 'Racing Bulls', 'Liberty Media', 'Formula 1', 'Formula One',
    'Apple', 'Netflix', 'ESPN', 'Sky', 'Comcast', 'Dorilton', 'Mumtalakat', 'Lawrence Stroll', 'Gene Haas',
    'Stellantis', 'Geely', 'Lenovo', 'LVMH', 'Aramco', 'Saudi Aramco', 'Pirelli', 'DHL', 'Heineken', 'Crypto.com',
    'Salesforce', 'AWS', 'Amazon', 'Qatar Airways', 'Emirates', 'MSC', 'Paramount Plus')}
JUNK = re.compile(r'freelanc|self[- ]?employed|no[- ]company|independiente|follow us|friends\.com|stealth|'
                  r'\bconfidential\b|\bunemployed\b|\bretired\b|\bjob seeker|\bmy company\b|\btest\b|\bdemo\b|^n/?a$', re.I)
GOVRX = re.compile(r'central bank|reserve bank|development bank|state bank of|\bministry\b|\bministerio\b|'
                   r'\bauthority\b|\bmunicipal|\bcounty\b|\bcity of\b|\bgovernment\b|\bfederal\b|\bdepartment of\b|'
                   r'\bpublic sector\b|\bazienda sanitaria|\bazienda ospedal|\bnhs\b|\bpolice\b|\barmy\b|\bnavy\b|'
                   r'\bair force\b|\bfannie mae\b|\bfreddie mac\b|\buniversit', re.I)

rej = []          # (record, reason)
keep = []
seen_known = 0
for k in order:
    r = recs[k]
    n = r['key']
    if (n and n in KN) or (r['domain'] and r['domain'] in KD):
        seen_known += 1; continue
    t = r['type']
    if n in TIES or any(len(s) >= 5 and (n == s or (n.startswith(s) and len(n) - len(s) <= 6)) for s in SP) or n in SP:
        rej.append((r, 'F1 sponsor / tie already on record')); continue
    if any(n == s or (len(s) >= 5 and n.startswith(s)) for s in F1_CORE):
        rej.append((r, 'F1 team / engine maker / rights holder or its group')); continue
    if JUNK.search(r['name']):
        rej.append((r, 'Suspect record (placeholder / freelancer / spam)')); continue
    if GOVRX.search(r['name']):
        rej.append((r, 'Government / public body')); continue
    if r['owner'] and (strict(r['owner']) in SP or strict(r['owner']) in TIES):
        rej.append((r, f"Parent is an F1 sponsor ({r['owner']})")); continue
    if any(s in n for s in SANCTION if len(s) >= 4) or n in SANCTION:
        rej.append((r, 'Possible sanctions-list name (to confirm)')); continue
    if t in ('Government Agency', 'Non Profit', 'Educational'):
        rej.append((r, f'{t} - not a commercial sponsor')); continue
    if r['owner'] and strict(r['owner']) in KN:
        rej.append((r, f"Subsidiary of a brand already in your files ({r['owner']})")); continue
    if r['country'].lower() in BADCOUNTRY or r['domain'].endswith(('.ru', '.by', '.ir')):
        rej.append((r, 'Sanctioned country (Russia / Belarus / Iran etc.)')); continue
    a = arm_of_known(r)
    if a:
        rej.append((r, f'Regional arm / unit of a brand already in your files ({a})')); continue
    if r['domain'].endswith('.gov') or '.gov.' in r['domain']:
        rej.append((r, 'Government body')); continue
    if 'Clay' in r['sources'] and r['emp'] and r['emp'] < 51:
        rej.append((r, f"Under 51 staff ({r['emp']})")); continue
    if t in ('Self Employed', 'Self Owned') and r['emp'] < 200:
        rej.append((r, 'Suspect record (self-employed / tiny)')); continue
    if not r['domain']:
        rej.append((r, 'Suspect record (no website)')); continue
    if r['score_rev'] >= 75 and 0 < r['emp'] < 100 and 'Apollo' not in r['sources']:
        rej.append((r, f"Suspect data (USD 1B+ revenue but {r['emp']} staff)")); continue
    code = sector_of(r)
    if code == 'GOV':
        rej.append((r, 'Government body')); continue
    r['sector'] = code
    keep.append(r)

# fold sibling records that share a website root (e.g. brand.com / brand.de) - keep the strongest
grp = defaultdict(list)
for r in keep:
    s = sld(r['domain'])
    if len(s) >= 6: grp[s].append(r)
drop = set()
for s, rs in grp.items():
    if len(rs) < 2: continue
    rs.sort(key=lambda r: (-r['score_rev'], -r['emp']))
    head = rs[0]
    for r in rs[1:]:
        drop.add(id(r)); rej.append((r, f"Duplicate / regional entity of {head['name']} (same website root)"))
keep = [r for r in keep if id(r) not in drop]


def base_of(name):
    t = toks(name)
    while len(t) > 1 and t[-1] in GEO: t.pop()
    return ''.join(t), len(t) < len(toks(name))


# collapse regional entities ("Samsung Switzerland", "Randstad Deutschland") onto one parent-brand row
bykey = {r['key']: r for r in keep}
grp = defaultdict(list)
for r in keep:
    b, regional = base_of(r['name'])
    if regional and len(b) >= 4:
        if b in KN:
            drop.add(id(r)); rej.append((r, f'Regional arm of a brand already in your files ({b})')); continue
        grp[b].append(r)
for b, rs in grp.items():
    head = bykey.get(b)
    if head is None or id(head) in drop:
        rs.sort(key=lambda r: (-r['score_rev'], -r['emp']))
        head = rs[0]; rs = rs[1:]
        head['flags_pre'] = 'regional entity - assess parent brand'
    for r in rs:
        if r is head: continue
        drop.add(id(r)); rej.append((r, f"Regional entity folded into {head['name']}"))
keep = [r for r in keep if id(r) not in drop]

# ---------------------------------------------------------------- scoring
for r in keep:
    b = r['rev_band']
    if r['score_rev'] >= 75: route = 'Revenue USD 1B+'
    elif b == '500M-1B': route = 'Revenue USD 500M-1B'
    elif b in ('200M-500M', '75M-200M'): route = 'Revenue USD 150-500M (Stretch cap)'
    else: route = ''
    if r['funded']:
        route = (route + ' + ' if route else '') + 'Raised USD 100M+'
    r['route'] = route
    flags = [r['flags_pre']] if r.get('flags_pre') else []
    if b == '75M-200M': flags.append('revenue may be under USD 150M')
    if r['owner']: flags.append(f"owned by {r['owner']}")
    if r['funded'] and not r['score_rev']: flags.append('up-and-coming (funded, revenue n/a)')
    if r['funded'] and 50 <= r['score_rev'] <= 75: flags.append('up-and-coming')
    r['flags'] = '; '.join(flags)
    base = r['score_rev'] or 55
    if r['funded']: base = max(base, 58) + 4
    emp_bonus = min(10.0, 2.0 * math.log10(max(r['emp'], 1)))
    pub = 3 if r['type'] == 'Public Company' or r['exchange'] else 0
    cap = 'Stretch cap' in route and not r['funded']
    r['score'] = round(min(base + emp_bonus + pub, 69 if cap else 999), 1)

# ---------------------------------------------------------------- workbook
HDR = ['rank', 'company', 'website', 'country', 'qualifying route', 'revenue band (database)', 'staff',
       'ownership', 'database industry', 'source', 'pre-score', 'flags', 'linkedin', 'status']
bold = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F3864')
wb = openpyxl.Workbook()
ws = wb.active; ws.title = '00 READ ME'
by = defaultdict(list)
for r in keep: by[r['sector']].append(r)
xtabs = sorted(k for k in by if k not in S)
lines = [
    'Wave 15 - database coverage sweep (UNCHECKED first cut)',
    '',
    'What this is: every company Clay and Apollo hold at USD 150M+ revenue or USD 100M+ funding,',
    'minus anything already in your master / Wave 14 / batch / beauty / expansion files and the F1 ties register.',
    'Sorted into your 56 master sectors; brands outside them sit on the X tabs. Ranked best-first by pre-score.',
    '',
    'Status = UNCHECKED: database data only. Revenue is often a Clay/Apollo estimate and sector placement is automatic.',
    'Every brand will be verified against the v4 rules; checked tabs will replace these in batches.',
    '',
    'Rules (v4): pass on ANY of - worth USD 1B+ / raised USD 100M+ in last 3 yrs / revenue USD 500M+ /',
    'revenue USD 150-500M (capped at Stretch). Then: 51+ staff, no distress, no F1 tie (direct or parent), not sanctioned.',
    '',
    'Pre-score: revenue band (100B+ = 100 ... 75-200M = 40), +funding, +staff size, +listed. Capped at 69 on the',
    'USD 150-500M route. This is only a starting order - the checked version uses the full v3/v4 priority score.',
    '',
    f'Pulled (unique): {len(order):,}    Already in your files: {seen_known:,}    Auto-removed: {len(rej):,}    New to check: {len(keep):,}',
    '',
    'Tab / count:',
]
for l in lines: ws.append([l])
for s in SECTORS: ws.append([s, len(by.get(s[:2], []))])
for x in xtabs: ws.append([x, len(by[x])])
ws.append(['99 AUTO-REMOVED', len(rej)])
ws.column_dimensions['A'].width = 110
ws['A1'].font = Font(bold=True, size=14)


def write(ws, rows):
    ws.append(HDR)
    for c in ws[1]: c.font = bold; c.fill = fill
    rows.sort(key=lambda r: (-r['score'], r['name'].lower()))
    for i, r in enumerate(rows, 1):
        ws.append([i, r['name'], r['domain'], r['country'], r['route'], r['rev_band'], r['emp'] or '',
                   r['type'] or ('Listed ' + r['exchange'] if r['exchange'] else ''), r['industry'] or ('NAICS ' + r['naics'] if r['naics'] else ''),
                   '+'.join(sorted(r['sources'])), r['score'], r['flags'], r['linkedin'], 'Unchecked'])
    for col, w in zip('ABCDEFGHIJKLMN', (6, 38, 28, 16, 34, 14, 9, 18, 30, 12, 9, 40, 40, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C2'


for s in SECTORS: write(wb.create_sheet(s[:31]), list(by.get(s[:2], [])))
for x in xtabs: write(wb.create_sheet(x[:31]), list(by[x]))
wr = wb.create_sheet('99 AUTO-REMOVED')
wr.append(['company', 'website', 'country', 'reason', 'database industry', 'revenue band', 'staff', 'source'])
for c in wr[1]: c.font = bold; c.fill = fill
for r, why in sorted(rej, key=lambda t: (t[1], t[0]['name'].lower())):
    wr.append([r['name'], r['domain'], r['country'], why, r['industry'], r['rev_band'], r['emp'] or '', '+'.join(sorted(r['sources']))])
for col, w in zip('ABCDEFGH', (38, 28, 16, 50, 30, 12, 8, 12)): wr.column_dimensions[col].width = w
wb.save(sys.argv[1])
json.dump([{k: (sorted(v) if isinstance(v, set) else v) for k, v in r.items()} for r in keep],
          open(f'{W}/w15_candidates.json', 'w'))
print('pulled', len(order), 'known', seen_known, 'removed', len(rej), 'keep', len(keep))
print(Counter(re.split(r' \(| of | into ', why)[0] for _, why in rej).most_common())
print(sorted(((len(by[k]), k) for k in by), reverse=True)[:70])
