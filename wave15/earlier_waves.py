"""Brands delivered in earlier waves (reports/*.xlsx, excluding Wave 15) that are NOT in the user's uploaded master.
Writes wave15/earlier_waves.json (one row per brand, best verdict kept). Usage: python3 -I wave15/earlier_waves.py"""
import glob, json, re, sys
import openpyxl
sys.path.insert(0, '/home/user/rg42/wave15')
W = '/home/user/rg42/wave15'
src = open(f'{W}/screen_merge.py').read()
ns = {}
exec(src[:src.index('KL = [')], ns)          # strict(), TIES, SP
strict, TIES, SP = ns['strict'], ns['TIES'], ns['SP']

up = openpyxl.load_workbook(f'{W}/inputs/MASTER_REFERENCE_OCT_upload.xlsx', read_only=True)
MN, MD = set(), set()
for ws in up.worksheets:
    for r in ws.iter_rows(values_only=True):
        for v in r:
            if not isinstance(v, str): continue
            v = v.strip()
            if re.fullmatch(r'[a-z0-9.-]+\.[a-z]{2,}', v.lower()): MD.add(v.lower().removeprefix('www.'))
            elif 1 < len(v) < 90: MN.add(strict(v))
MN.discard('')
SKIP = re.compile(r'read ?me|reject|dropped|pending|flag|summary|index|^00|^98|^99|changes|send|log|priority list|overview|notes', re.I)
VR = {'yes': 0, 'stretch': 1, 'low-value': 2, 'low value': 2, '': 3}
best = {}
for f in sorted(glob.glob('/home/user/rg42/reports/*.xlsx')):
    if 'Wave15' in f or 'MASTER-REFERENCE-OCT' in f: continue
    fn = f.split('/')[-1].replace('.xlsx', '')
    wb = openpyxl.load_workbook(f, read_only=True)
    for ws in wb.worksheets:
        if SKIP.search(ws.title): continue
        it = ws.iter_rows(values_only=True)
        h = next(it, None)
        if not h: continue
        h = [str(x or '').strip().lower() for x in h]
        def col(*names):
            for n in names:
                for i, x in enumerate(h):
                    if x == n or x.startswith(n): return i
        ci = col('company', 'brand', 'key_brand')
        if ci is None: continue
        di, vi, si, pi = col('domain'), col('verdict', 'tier'), col('sector_tab', 'sector'), col('priority_score', 'priority score')
        for r in it:
            if ci >= len(r) or not r[ci] or not isinstance(r[ci], str): continue
            name = r[ci].strip(); k = strict(name)
            if len(k) < 2: continue
            dom = str(r[di]).strip().lower().removeprefix('www.') if di is not None and di < len(r) and r[di] else ''
            if k in MN or (dom and dom in MD) or k in TIES or k in SP: continue
            v = str(r[vi]).strip() if vi is not None and vi < len(r) and r[vi] else ''
            vl = v.lower()
            if vl.startswith('no') or 'tie' in vl or 'reject' in vl: continue
            vv = 'Yes' if vl.startswith('yes') else 'Stretch' if vl.startswith('stretch') else 'Low-value' if vl.startswith('low') else ''
            sc = r[pi] if pi is not None and pi < len(r) and isinstance(r[pi], (int, float)) else None
            sec = str(r[si]) if si is not None and si < len(r) and r[si] else ws.title
            row = dict(name=name, domain=dom, verdict=vv, score=sc, sector=sec[:40], source=f'{fn} | {ws.title}')
            o = best.get(k)
            if o is None or (VR[vv.lower()], -(sc or 0)) < (VR[o['verdict'].lower()], -(o['score'] or 0)):
                if o: row['also_in'] = o['source']
                best[k] = row
rows = sorted(best.values(), key=lambda r: (VR[r['verdict'].lower()], -(r['score'] or 0), r['name'].lower()))
json.dump(rows, open(f'{W}/earlier_waves.json', 'w'), ensure_ascii=False, indent=0)
from collections import Counter
print('brands not in uploaded master', len(rows), Counter(r['verdict'] or '(no verdict)' for r in rows))
print('boom?', [r for r in rows if 'boom' in r['name'].lower()][:2])
