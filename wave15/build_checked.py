"""Build the CHECKED Wave 15 workbook from screening outputs (and stage-2 results when present).
Usage: python3 -I wave15/build_checked.py <out.xlsx>   (run wave15/screen_merge.py first)"""
import glob, json, re, sys
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill

W = '/home/user/rg42/wave15'
sys.path.insert(0, W)
SECTORS = ['01 Banking & Insurance', '02 Energy & Industrial', '03 Luxury & Beauty', '04 Food & Beverage',
           '05 Supercomputing', '06 UCaaS & CPaaS', '07 Marketing & Sales', '08 Devtools & Integration',
           '09 Fintech', '10 ERP & Finance Procurement', '11 Storage Hardware', '12 Trading & Crypto',
           '13 AI Agents', '14 Cyber + IT', '15 Observability', '16 Travel & Hospitality',
           '17 AI Cloud & Data', '18 Services & Other', '19 Gulf Sovereign & Tourism', '20 Staffing & Talent',
           '21 Materials & Mech Parts', '22 SaaS', '23 Machine Tools', '24 Logistics',
           '25 IT Services & Consulting', '26 Horizontal SaaS', '27 Chip Makers', '28 AI Data Services',
           '29 Tech Media & Gaming', '30 Commercial Vehicles', '31 GRC & Risk SW', '32 Defense & Industrial AI',
           '33 Automation & Robotics', '34 DC Power & Cooling', '35 Semiconductor Equipment', '36 IoT & Connectivity',
           '37 EDA', '38 Edge & CDN', '39 Construction & Field Svc', '40 Test & Measurement',
           '41 OT & Physical Security', '42 Health & Wellness', '43 Mobility & Automotive', '44 Grid & Transformers',
           '45 Electronic Parts', '46 Tools & Industrial Supply', '47 Vertical SaaS', '48 Telecom & Space',
           '49 Home & Living', '50 Law Firms & Legal AI', '51 Industrial Gases', '52 Photonics & Optical',
           '53 Consumer & Lifestyle', '54 TIC', '55 Packaging & Food Machinery', '56 Robotics AI']
S = {s[:2]: s for s in SECTORS}


# Wave 15 broad X tabs -> Wave 14 sector tabs
XMAP = {'pharma & biotech': 'NEW Pharma & Biotech', 'medtech': 'NEW Medtech & Healthcare Services',
        'retail & e-commerce': 'X06 Retail & E-commerce', 'real estate': 'X02 Real Estate & Building Products',
        'utilities & power': 'NEW Utilities & Power', 'mining & metals': 'X04 Mining, Metals & Commodities',
        'asset mgmt & pe': 'X01 Asset Mgmt, PE & Trading Firms', 'aviation & aerospace': 'X07 Aviation & Aerospace Suppliers',
        'education': 'NEW Education & EdTech', 'gambling & lotteries': 'NEW Gambling & Lotteries',
        'sports clubs & leagues': 'NEW Sports Clubs & Leagues', 'maritime & shipbuilding': 'X09 Maritime, Shipping & Ports',
        'rail & transit': 'X12 Rail, Transit & Road Infra', 'consumer electronics': 'X15 Consumer Electronics & Audio',
        'agriculture': 'X05 Agriculture & AgTech', 'waste & recycling': 'X30 Waste, Recycling & Circular',
        'live events & ticketing': 'X24 Ticketing, Live Events & Experiences', 'pet care': 'X19 Pet Care & Nutrition',
        'other': 'NEW Other'}


def tab_of(sec, fallback):
    s = str(sec or '').strip()
    m = re.match(r'^(X\d{2}|NEW)\b', s, flags=re.I)
    if m: return (m.group(1).upper() + s[len(m.group(1)):]).strip()[:31]
    if s.upper().startswith('X'):
        k = re.sub(r'^X[\s\-:]*', '', s, flags=re.I).strip().lower()
        if k in XMAP: return XMAP[k][:31]
    m = re.match(r'^(\d{1,2})\b', s)
    if m and m.group(1).zfill(2) in S: return S[m.group(1).zfill(2)]
    if s.upper().startswith('X'):
        return ('X ' + re.sub(r'^X[\s\-:]*', '', s, flags=re.I)).strip()[:31]
    f = str(fallback or '')
    return S.get(f[:2], f[:31] or 'X Other')


def prio(v, fit):
    try: fit = int(fit)
    except (TypeError, ValueError): fit = 0
    return {'Yes': 60 + 8 * fit, 'Stretch': 40 + 6 * fit, 'Low-value': 20 + 4 * fit}.get(v, 0)


kept = json.load(open(f'{W}/screen_kept.json'))
deep = {}
for f in glob.glob(f'{W}/verify/out_*.json'):
    try:
        for r in json.load(open(f)): deep[r.get('id') or r.get('company')] = r
    except Exception:
        pass
cands = json.load(open(f'{W}/w15_candidates.json'))
screened_ids = set()
drops = []
for f in sorted(glob.glob(f'{W}/screen/out_*.json')):
    for r in json.load(open(f)):
        screened_ids.add(r.get('id'))
        if not (r.get('keep') is True or str(r.get('keep')).lower() == 'true'):
            drops.append(r)

rows = []
for r in kept:
    d = deep.get(r.get('id')) or {}
    v = d.get('verdict') or r.get('verdict_guess') or ''
    if v == 'No': continue
    fit = d.get('fit_score', r.get('fit'))
    rows.append(dict(
        tab=tab_of(d.get('sector_tab') or r.get('sector'), r.get('sector_guess')), name=r['name'], verdict=v,
        score=prio(v, fit), fit=fit, route=r.get('scale_route', ''), evidence=d.get('valuation') or r.get('scale_evidence', ''),
        staff=d.get('headcount_band') or r.get('staff_est', ''), tie=r.get('f1_tie', ''), parent=r.get('parent', ''),
        status=r.get('status', ''), web=r.get('website', ''), country=r.get('country', ''),
        note=d.get('verdict_reason') or r.get('note', ''), stage='Fully verified' if d else 'Checked (stage 1)',
        sells=d.get('sells', ''), angle=d.get('best_angle', ''), hq=d.get('hq', ''),
        upc='yes' if d.get('up_and_coming') is True else '',
        flags='; '.join(d.get('flags') or []) if isinstance(d.get('flags'), list) else str(d.get('flags') or ''),
        src=' '.join((d.get('sources') or [])[:3]) if isinstance(d.get('sources'), list) else str(d.get('sources') or '')))

HDR = ['rank', 'company', 'verdict', 'priority score', 'fit (0-4)', 'qualifying route', 'scale evidence', 'staff',
       'F1 tie check', 'parent / owner', 'status', 'website', 'country', 'notes', 'check stage', 'also on',
       'what they sell', 'best F1 angle', 'HQ', 'up-and-coming', 'flags', 'sources']
bold = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F3864')
red = PatternFill('solid', fgColor='F8CBAD')
VORD = {'Yes': 0, 'Stretch': 1, 'Low-value': 2}


def write(ws, rs, also=True):
    ws.append(HDR)
    for c in ws[1]: c.font = bold; c.fill = fill
    rs = sorted(rs, key=lambda r: (-r['score'], VORD.get(r['verdict'], 9), r['name'].lower()))
    for i, r in enumerate(rs, 1):
        ws.append([i, r['name'], r['verdict'], r['score'], r['fit'], r['route'], r['evidence'], r['staff'], r['tie'],
                   r['parent'], r['status'], r['web'], r['country'], r['note'], r['stage'],
                   ('01 PRIORITY' if (also and r['verdict'] == 'Yes') else ''),
                   r['sells'], r['angle'], r['hq'], r['upc'], r['flags'], r['src']])
        if also and r['verdict'] == 'Yes':
            ws.cell(row=i + 1, column=2).fill = red
    for col, w in zip('ABCDEFGHIJKLMNOPQRSTUV', (6, 34, 10, 8, 7, 22, 40, 12, 34, 24, 14, 24, 14, 50, 16, 12, 40, 40, 18, 10, 40, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'C2'


by = defaultdict(list)
for r in rows: by[r['tab']].append(r)
wb = openpyxl.Workbook()
ws = wb.active; ws.title = '00 READ ME'
cnt = Counter(r['verdict'] for r in rows)
todo = len([c for c in cands if (c['domain'] or c['key']) not in screened_ids])
for l in ['Wave 15 - database sweep, CHECKED brands (new names only)', '',
          'Every brand here came from Clay / Apollo, is NOT in your master or other files, and has been checked',
          'against the v4 rules (real & operating, independent brand, scale route, 51+ staff, no distress,',
          'no F1 tie direct or via parent, not sanctioned). Ranked by priority score: Yes 60+8xfit, Stretch 40+6xfit,',
          'Low-value 20+4xfit. "check stage" says whether the brand has had the full sourced check yet.', '',
          'Qualifying routes: worth USD 1B+ / raised USD 100M+ in 36 months / revenue USD 500M+ /',
          'revenue USD 150-500M (capped at Stretch).', '',
          f"Checked so far: {len(screened_ids):,} brands -> kept {len(rows):,} (Yes {cnt['Yes']}, Stretch {cnt['Stretch']}, "
          f"Low-value {cnt['Low-value']}); dropped {len(drops):,} (reasons on tab 99).",
          f'Still to check: {todo:,} database names.',
          'Sources: Clay (USD 150M+ revenue, 13 industry groups) + Apollo (USD 100M+ funded) + parent companies of',
          'dropped subsidiaries. Some USD 200M-1B industry pulls were capped at 4,000 rows each by the database.',
          'Rows whose notes say "(est)" rest on estimated figures; confirm before outreach.', '', 'Tab / count:']:
    ws.append([l])
ws['A1'].font = Font(bold=True, size=14)
ws.column_dimensions['A'].width = 110
pr = [r for r in rows if r['verdict'] == 'Yes']
ws.append(['01 PRIORITY (all Yes)', len(pr)])
for s in SECTORS: ws.append([s, len(by.get(s, []))])
xt = sorted(k for k in by if k not in SECTORS)
for x in xt: ws.append([x, len(by[x])])
ws.append(['99 DROPPED', len(drops)])
write(wb.create_sheet('01 PRIORITY'), pr, also=False)
for s in SECTORS: write(wb.create_sheet(s[:31]), by.get(s, []))
for x in xt: write(wb.create_sheet(x[:31]), by[x])
wd = wb.create_sheet('99 DROPPED')
wd.append(['company', 'reason', 'parent', 'status', 'F1 tie check', 'website'])
for c in wd[1]: c.font = bold; c.fill = fill
for r in sorted(drops, key=lambda r: str(r.get('reason', ''))):
    wd.append([r.get('real_name') or r.get('company'), r.get('reason', ''), r.get('parent', ''), r.get('status', ''),
               r.get('f1_tie', ''), r.get('id', '')])
for col, w in zip('ABCDEF', (34, 50, 30, 16, 40, 28)): wd.column_dimensions[col].width = w
wb.save(sys.argv[1])
print('rows', len(rows), dict(cnt), 'drops', len(drops), 'todo', todo, 'tabs', len(wb.sheetnames))
