"""Write the partner database (data.json) to Excel: all partners, team summary, and brands checked but excluded."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
D = json.load(open(sys.argv[1]))
ROW = {1: "Title / principal", 2: "Official / team partner", 3: "Supplier / regional / licensee"}
SEC = {"FS": "Financial services", "PS": "Professional services", "LS": "Life sciences & health", "O": "Other"}
wb = Workbook()
ws = wb.active; ws.title = "Partners"
ws.append(["Team", "Group", "Partner", "Tier (as published)", "Tier (harmonised)", "Wolters Kluwer focus sector",
           "Status", "Announced in 2026", "Source"])
for t in D["teams"]:
    for p in t["partners"]:
        ws.append([t["name"], t["group"], p["name"], p["tier"], ROW[p["row"]], SEC[p["sector"]],
                   f'Signed, starts {p["starts"]}' if p["starts"] else "Active 2026",
                   f'{p["new"]} 2026' if p["new"] else "", D["sources"][t["key"]]])
s2 = wb.create_sheet("Summary")
s2.append(["Team", "Group", "Active partners", "Signed from 2027", "Title / principal", "Official / team",
           "Supplier / regional / licensee", "Financial services", "Professional services", "Life sciences & health",
           "Focus-sector share", "Announced in 2026"])
for t in D["teams"]:
    cur = [x for x in t["partners"] if not x["starts"]]
    s2.append([t["name"], t["group"], len(cur), len(t["partners"]) - len(cur)] + [sum(x["row"] == r for x in cur) for r in (1, 2, 3)]
              + [sum(x["sector"] == k for x in cur) for k in ("FS", "PS", "LS")]
              + [sum(x["sector"] != "O" for x in cur) / len(cur), sum(1 for x in t["partners"] if x["new"])])
s3 = wb.create_sheet("Checked, not included")
s3.append(["Reason", "Brands"])
for reason, brands in D["excluded"]:
    s3.append([reason, brands])
for sh in (ws, s2, s3):
    for c in sh[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="C00000")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    sh.freeze_panes = "A2"; sh.auto_filter.ref = sh.dimensions
    for i, col in enumerate(sh.columns, 1):
        sh.column_dimensions[get_column_letter(i)].width = min(70, max(12, max(len(str(c.value or "")) for c in col) + 2))
for row in s3.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
for c in s2["K"][1:]: c.number_format = "0%"
wb.save(sys.argv[2])
print(ws.max_row - 1, "partner rows")
