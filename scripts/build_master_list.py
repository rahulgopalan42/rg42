#!/usr/bin/env python3
"""Build the final F1 prospect list of companies with $100M+ total funding.

Combines the Clay pulls (private and all other company types), the per-company
clean-up review, and companies added from 2025-26 funding-round news, then writes
an Excel workbook (Summary, Private & other, Public, Removed) and a CSV.

Placement rules, from the review:
- avoid / f1_partner / subsidiary with high confidence  -> Removed tab, with the reason
- listed with high confidence                           -> Public tab
- keep with high confidence                             -> Private & other tab
- anything with low confidence                          -> stays where Clay's company type puts it,
                                                           with a "Check" note

Usage: build_master_list.py <scratch_dir> <output_base>
"""
import csv
import glob
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import categorize_sectors as cs  # noqa: E402

PRIVATE_TYPES = {"Privately Held", "Self Owned"}
REMOVE_REASON = {"avoid": "Bankrupt, defunct, sanctions risk or bad record",
                 "f1_partner": "Already an F1 partner", "subsidiary": "Subsidiary of a larger group"}
# Clay entityIds of companies that own an F1 team or Formula 1 itself (not sponsor prospects)
F1_OWNERS = {"22114834": "Ferrari F1 team", "2684047": "Formula 1 (Liberty Media)",
             "36598105": "Mercedes F1 team (co-owner)"}
STATUS_WORDS = {"listed": "listed on a stock exchange", "subsidiary": "a subsidiary",
                "f1_partner": "already an F1 partner", "avoid": "one to avoid", "keep": "not listed"}


def norm_domain(d):
    d = re.sub(r"^https?://", "", (d or "").lower().strip()).split("/")[0].split(".")
    while d and d[0] in ("www", "recrute", "careers", "jobs", "life", "international"):
        d = d[1:]
    return ".".join(d)


def norm_name(n):
    n = re.sub(r"[^a-z0-9 ]", " ", (n or "").lower())
    n = re.sub(r"\b(inc|llc|ltd|limited|plc|corp|corporation|co|company|group|holdings?|s a|sa|ag|gmbh|the)\b", " ", n)
    return " ".join(n.split())


def read_csvs(pattern):
    rows = []
    for p in sorted(glob.glob(pattern)):
        rows.extend(csv.DictReader(open(p, newline="")))
    return rows


def main(scratch, out_base):
    clay = {}
    for r in read_csvs(os.path.join(scratch, "clay_pull", "seg_*.csv")) + \
            read_csvs(os.path.join(scratch, "clay_pull_other", "*.csv")):
        clay.setdefault(r["entityId"], r)
    review = {r["entityId"]: r for r in read_csvs(os.path.join(scratch, "review", "*_review_*.csv"))}
    missing_review = [e for e in clay if e not in review]

    rows = []
    for eid, r in clay.items():
        rv = review.get(eid, {"status": "keep", "confidence": "low", "detail": "", "note": "not reviewed"})
        status, conf = rv["status"], rv["confidence"]
        detail, note = rv.get("detail", ""), rv.get("note", "")
        clay_private = r["type"] in PRIVATE_TYPES
        check = ""
        if eid in F1_OWNERS:
            tab, reason, detail = "Removed", "Owns an F1 team or Formula 1", F1_OWNERS[eid]
        elif conf == "high" and status in REMOVE_REASON:
            tab, reason = "Removed", REMOVE_REASON[status]
        elif conf == "high" and status == "listed":
            tab, reason = "Public", ""
        elif conf == "high":
            tab, reason = "Private & other", ""
        else:
            tab, reason = ("Public" if r["type"] == "Public Company" else "Private & other"), ""
            if status != "keep":
                check = f"Check: possibly {STATUS_WORDS[status]}" + (f" ({detail})" if detail else "")
        sector, f1 = cs.classify(r)
        rows.append({
            "entityId": eid, "name": r["name"].strip(), "tab": tab, "removed_reason": reason,
            "review_status": status, "review_detail": detail, "review_note": note, "check": check,
            "confidence": conf, "sector": sector, "f1_category": f1, "industry": r["industry"],
            "funding_range": r["funding_range"], "clay_type": r["type"] or "(none)", "country": r["country"],
            "city": r["locality"], "size": r["size"], "employees": r["employee_count"],
            "revenue": r["annual_revenue"], "domain": r["domain"], "linkedin": r["linkedin_url"],
            "source": "Clay", "round": "", "round_amount_m": "", "announced": "", "news_url": "",
        })

    # Collapse duplicate Clay profiles of the same company (same domain and same normalised name)
    seen, deduped, dupes = {}, [], 0
    for r in sorted(rows, key=lambda r: (r["tab"] == "Removed", int(r["entityId"]))):
        key = (norm_domain(r["domain"]), norm_name(r["name"]))
        if key[0] and key in seen:
            dupes += 1
            continue
        seen[key] = r
        deduped.append(r)
    rows = deduped

    # Funding-news rounds: enrich matching Clay rows, add the rest as new companies
    news = read_csvs(os.path.join(scratch, "news_pull", "rounds_*.csv"))
    news_review = {norm_name(r["company"]): r for r in read_csvs(os.path.join(scratch, "news_pull", "news_review*.csv"))}
    by_domain = {norm_domain(r["domain"]): r for r in rows if r["domain"]}
    by_name = defaultdict(list)
    for r in rows:
        by_name[norm_name(r["name"])].append(r)
    best_round = {}
    for n in news:
        try:
            amt = float(n["amount_usd_m"])
        except ValueError:
            continue
        if amt < 100:
            continue
        k = norm_name(n["company"])
        if k not in best_round or float(best_round[k]["amount_usd_m"]) < amt:
            best_round[k] = n
    by_eid = {r["entityId"]: r for r in rows}
    added, matched = 0, 0
    for k, n in best_round.items():
        rv = news_review.get(k, {})
        if rv.get("duplicate_of"):
            continue  # same company as another news row, which carries it
        target = by_domain.get(norm_domain(n["website"])) if n["website"] else None
        if target is None and len(by_name.get(k, [])) == 1:
            target = by_name[k][0]
        if target is None and rv.get("clay_entityId"):
            target = by_eid.get(rv["clay_entityId"].strip())
        fields = {"round": n["round_type"], "round_amount_m": n["amount_usd_m"], "announced": n["announced"],
                  "news_url": n["source_url"]}
        if target is not None:
            target.update(fields)
            matched += 1
            continue
        status, conf = rv.get("status", "keep"), rv.get("confidence", "low")
        sector = rv.get("sector") or cs.OTHER
        if conf == "high" and status in REMOVE_REASON:
            tab, reason = "Removed", REMOVE_REASON[status]
        elif status == "listed" and conf == "high":
            tab, reason = "Public", ""
        else:
            tab, reason = "Private & other", ""
        amt = float(n["amount_usd_m"])
        rows.append({
            "entityId": "", "name": n["company"].strip(), "tab": tab, "removed_reason": reason,
            "review_status": status, "review_detail": rv.get("detail", ""), "review_note": rv.get("note", ""),
            "check": ("Check: possibly " + STATUS_WORDS[status]) if conf == "low" and status != "keep" else "",
            "confidence": conf, "sector": sector, "f1_category": cs.SECTOR_F1.get(sector, cs.NOT_IN_F1),
            "industry": n["sector_hint"], "funding_range": "$250M+" if amt >= 250 else "$100M - $250M",
            "clay_type": "(not in Clay)", "country": n["country"], "city": "", "size": "", "employees": "",
            "revenue": "", "domain": norm_domain(n["website"]), "linkedin": "", "source": "Funding news",
            **fields,
        })
        added += 1

    order = {s: i for i, s in enumerate(cs.SECTOR_ORDER)}
    fund = {"$250M+": 0, "$100M - $250M": 1}
    rows.sort(key=lambda r: (order.get(r["sector"], 99), fund.get(r["funding_range"], 9), r["name"].lower()))

    fields = ["name", "tab", "sector", "f1_category", "funding_range", "round", "round_amount_m", "announced",
              "news_url", "source", "check", "removed_reason", "review_detail", "review_note", "confidence",
              "industry", "clay_type", "country", "city", "size", "employees", "revenue", "domain", "linkedin",
              "entityId"]
    with open(out_base + ".csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    headers = ["Company", "Sector", "F1 sponsorship category", "Total funding (range)", "Largest 2025-26 round",
               "Round ($M)", "Announced", "Round source", "Source", "Check", "Clay industry", "Clay company type",
               "Country", "City", "Size", "Employees", "Annual revenue", "Website", "LinkedIn"]
    widths = [34, 30, 26, 18, 22, 10, 12, 40, 13, 44, 30, 18, 18, 20, 13, 11, 15, 26, 44]

    def values(r):
        emp = r["employees"]
        return [r["name"], r["sector"], r["f1_category"], r["funding_range"], r["round"],
                float(r["round_amount_m"]) if r["round_amount_m"] else "", r["announced"], r["news_url"],
                r["source"], r["check"], r["industry"], r["clay_type"], r["country"], r["city"], r["size"],
                int(emp) if str(emp).isdigit() else emp, r["revenue"], r["domain"], r["linkedin"]]

    def sheet(ws, sheet_rows, extra=None):
        hdr = headers + ([extra[0]] if extra else [])
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in sheet_rows:
            ws.append(values(r) + ([extra[1](r)] if extra else []))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, wdt in enumerate(widths + ([50] if extra else []), start=1):
            ws.column_dimensions[get_column_letter(i)].width = wdt

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    tabs = ["Private & other", "Public"]
    by = {t: [r for r in rows if r["tab"] == t] for t in tabs + ["Removed"]}
    summary.append(["Sector"] + tabs + ["Total"])
    for c in summary[1]:
        c.font = Font(bold=True)
    for s in cs.SECTOR_ORDER:
        counts = [sum(r["sector"] == s for r in by[t]) for t in tabs]
        if sum(counts):
            summary.append([s] + counts + [sum(counts)])
    totals = [len(by[t]) for t in tabs]
    summary.append(["Total"] + totals + [sum(totals)])
    summary[summary.max_row][0].font = Font(bold=True)
    summary.append([])
    summary.append(["Removed", len(by["Removed"])])
    for reason, n in Counter(r["removed_reason"] for r in by["Removed"]).most_common():
        summary.append(["  " + reason, n])
    summary.append([])
    summary.append(["Added from 2025-26 funding news (not found in Clay)",
                    sum(r["source"] == "Funding news" and r["tab"] != "Removed" for r in rows)])
    summary.append(["Clay companies with a 2025-26 $100M+ round on record", matched])
    summary.append(["Rows with a Check note (uncertain review call)",
                    sum(bool(r["check"]) for r in rows if r["tab"] != "Removed")])
    summary.append([])
    summary.append(["Notes"])
    summary[summary.max_row][0].font = Font(bold=True)
    for note in [
        "Clay: every company with total funding of $100M+ (Clay's $100M-$250M and $250M+ ranges), all company types.",
        "\"$250M+\" means $250M and above, with no upper limit. Clay does not split it further.",
        "Clay misses some companies (it shows 'Funding unknown'); companies with a $100M+ round in 2025-26 funding news were added.",
        "Removed: subsidiaries, current F1 partners (SPORTFIVE playbook list), and bankrupt or sanctions-risk companies.",
        "Companies tagged private in Clay but actually listed were moved to the Public tab.",
        "Review calls are from public knowledge, not live data. Uncertain calls stay in the list with a Check note.",
    ]:
        summary.append([note])
    summary.column_dimensions["A"].width = 60
    for col in "BCD":
        summary.column_dimensions[col].width = 16

    for t in tabs:
        sheet(wb.create_sheet(t), by[t])
    sheet(wb.create_sheet("Removed"), by["Removed"],
          ("Why removed", lambda r: r["removed_reason"] + (f": {r['review_detail']}" if r["review_detail"] else "")))
    wb.save(out_base + ".xlsx")

    print("clay rows:", len(clay), "| missing review:", len(missing_review), "| duplicate profiles dropped:", dupes)
    print("news rounds >= $100M:", len(best_round), "| matched to Clay:", matched, "| added:", added)
    for t in tabs + ["Removed"]:
        print(f"{t}: {len(by[t])}")
    print("removed by reason:", dict(Counter(r["removed_reason"] for r in by["Removed"])))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
