#!/usr/bin/env python3
"""Pull every company / brand name and website out of reference workbooks.

Each sheet's header row is found by looking for columns such as "company", "Company name",
"Parent Company Name", "domain" or "Website". Name columns are the explicit name/brand
headers plus the column just left of "domain" (some tabs label it "x"). Sheets with no
header (e.g. a plain list in column A) contribute their short text cells.

Names like "Aubert & Duval (Eramet Group)" or "Kimberly-Clark (Huggies, Kleenex)" also
yield the names inside the brackets, marked as derived.

Usage: extract_reference_names.py <out_csv> <label>=<workbook.xlsx> [...]
"""
import csv
import re
import sys

from openpyxl import load_workbook

NAME_HEADERS = {
    "company", "company name", "parent company name", "ultimate parent company",
    "most relevant brand / subsidiary to approach", "most relevant subsidiary / brand", "brand", "brand name",
}
LIST_HEADERS = {"existing subsidiary / brand(s) in datasets"}
DOMAIN_HEADERS = {"domain", "website"}
SKIP_VALUES = {"x", "yes", "no", "key", "tbr", "tes", "none", "n/a", "na", "-", "?"}


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def split_brackets(name):
    """'A (B, C / D)' -> main 'A', derived ['B', 'C', 'D']."""
    main = re.sub(r"\s*\(.*?\)\s*", " ", name).strip(" -–,;")
    inner = []
    for grp in re.findall(r"\((.*?)\)", name):
        for part in re.split(r"[,;/]| and | & (?=[A-Z])|\bincl\.?\b|\bformerly\b|\bnow\b", grp):
            part = part.strip(" .:-–")
            if 2 <= len(part) <= 60 and not re.search(r"\d{4}|subsidiary|division|segment|brands? post|owned|backed|spin", part, re.I):
                inner.append(part)
    return main, inner


def header_row(rows):
    for i, r in enumerate(rows[:8]):
        low = [clean(v).lower() for v in r]
        if any(h in NAME_HEADERS | DOMAIN_HEADERS for h in low):
            return i, low
    return None, None


def main(out_path, sources):
    out = []
    for label, path in sources:
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            hi, hdr = header_row(rows)
            if hi is None:
                for ri, r in enumerate(rows):
                    for ci, v in enumerate(r[:3]):
                        s = clean(v)
                        if 2 <= len(s) <= 80 and s.lower() not in SKIP_VALUES and not re.fullmatch(r"[\d.,%$ -]+", s):
                            out.append((label, ws.title, ri + 1, "list", s, ""))
                continue
            name_cols = [i for i, h in enumerate(hdr) if h in NAME_HEADERS]
            list_cols = [i for i, h in enumerate(hdr) if h in LIST_HEADERS]
            dom_cols = [i for i, h in enumerate(hdr) if h in DOMAIN_HEADERS]
            for d in dom_cols:
                if d > 0 and d - 1 not in name_cols and hdr[d - 1] not in DOMAIN_HEADERS:
                    name_cols.append(d - 1)
            for ri, r in enumerate(rows[hi + 1:], start=hi + 2):
                dom = next((clean(r[d]) for d in dom_cols if d < len(r) and clean(r[d])), "")
                for c in name_cols:
                    s = clean(r[c]) if c < len(r) else ""
                    if not s or s.lower() in SKIP_VALUES or len(s) > 120 or re.fullmatch(r"[\d.,%$ -]+", s):
                        continue
                    mainname, inner = split_brackets(s)
                    out.append((label, ws.title, ri, "name", s, dom))
                    if mainname and mainname != s:
                        out.append((label, ws.title, ri, "name-main", mainname, dom))
                    for x in inner:
                        out.append((label, ws.title, ri, "derived", x, ""))
                for c in list_cols:
                    s = clean(r[c]) if c < len(r) else ""
                    for part in s.split(";"):
                        nm = re.sub(r"\s*\(.*$", "", part).strip()
                        if 2 <= len(nm) <= 80:
                            out.append((label, ws.title, ri, "derived", nm, ""))
                if dom and not any(clean(r[c]) for c in name_cols if c < len(r)):
                    out.append((label, ws.title, ri, "domain-only", "", dom))
    with open(out_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["workbook", "sheet", "row", "kind", "name", "domain"])
        w.writerows(out)
    print(f"{len(out)} reference entries written to {out_path}")


if __name__ == "__main__":
    srcs = [a.split("=", 1) for a in sys.argv[2:]]
    main(sys.argv[1], srcs)
